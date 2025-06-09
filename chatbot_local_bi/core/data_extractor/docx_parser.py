import os
# For DOCX parsing, the python-docx library is used.
# Add "python-docx" to your requirements.txt
try:
    import docx
except ImportError:
    print("python-docx not installed. DOCX parsing will not be available. "
          "Please install it with: pip install python-docx")

class DocxParser:
    def __init__(self):
        """
        Initializes the DocxParser.
        Checks if the required library (python-docx) is available.
        """
        if 'docx' not in globals():
            raise ImportError(
                "python-docx is not installed. "
                "DOCX parsing functionality is unavailable. "
                "Install it with 'pip install python-docx'."
            )

    def parse(self, file_path: str) -> str:
        """
        Parses a DOCX file and extracts all text content from paragraphs.
        This version does not extract text from tables, headers, or footers.

        Args:
            file_path (str): The path to the DOCX file.

        Returns:
            str: The extracted text content from the DOCX.
                 Returns an error message string if parsing fails.

        Raises:
            FileNotFoundError: If the file_path does not exist.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"DOCX file not found: {file_path}")

        try:
            doc = docx.Document(file_path)
            text_content = []
            for para in doc.paragraphs:
                text_content.append(para.text)

            return "\n".join(text_content)
        except Exception as e:
            # Log the error or handle it more gracefully
            print(f"Error parsing DOCX file '{file_path}': {e}")
            return f"Error: Could not parse DOCX file. Details: {e}"

    def parse_with_tables(self, file_path: str) -> tuple[str, list[list[list[str]]]]:
        """
        Parses a DOCX file and extracts text from paragraphs and tables.

        Args:
            file_path (str): The path to the DOCX file.

        Returns:
            tuple[str, list[list[list[str]]]]:
                - The extracted paragraph text content.
                - A list of tables, where each table is a list of rows,
                  and each row is a list of cell strings.
                Returns (error_message, []) if parsing fails.

        Raises:
            FileNotFoundError: If the file_path does not exist.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"DOCX file not found: {file_path}")

        try:
            doc = docx.Document(file_path)

            # Extract paragraph text
            paragraph_text = [para.text for para in doc.paragraphs]

            # Extract table data
            tables_data = []
            for table in doc.tables:
                table_content = []
                for row in table.rows:
                    row_content = [cell.text for cell in row.cells]
                    table_content.append(row_content)
                tables_data.append(table_content)

            return "\n".join(paragraph_text), tables_data
        except Exception as e:
            print(f"Error parsing DOCX file with tables '{file_path}': {e}")
            return f"Error: Could not parse DOCX file. Details: {e}", []

# Example Usage (for testing)
if __name__ == '__main__':
    # This example assumes you have a DOCX file named 'sample.docx'
    # in the same directory or provide a full path.
    # You would need to create a dummy DOCX for this test to run.

    parser = DocxParser()
    sample_docx_path = "test_sample.docx" # Replace with path to an actual DOCX

    # Create a dummy DOCX for testing if python-docx is available
    if 'docx' in globals():
        try:
            doc = docx.Document()
            doc.add_heading('Test Document Title', level=1)
            doc.add_paragraph('This is the first paragraph of the test document.')
            doc.add_paragraph('This is another paragraph with some more text.')

            # Add a table
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = 'Foo'
            table.cell(0, 1).text = 'Bar'
            table.cell(1, 0).text = 'Baz'
            table.cell(1, 1).text = 'Qux'

            doc.save(sample_docx_path)
            print(f"Created dummy DOCX: {sample_docx_path}")

            print("\n--- Parsing DOCX (paragraphs only) ---")
            content = parser.parse(sample_docx_path)
            if content.startswith("Error:"):
                print(content)
            else:
                print(f"Extracted Content:\n{content}")

            print("\n--- Parsing DOCX (paragraphs and tables) ---")
            para_content, tables = parser.parse_with_tables(sample_docx_path)
            if para_content.startswith("Error:"):
                print(para_content)
            else:
                print(f"Paragraph Content:\n{para_content}")
                print("\nTables Content:")
                if tables:
                    for i, table_data in enumerate(tables):
                        print(f"Table {i+1}:")
                        for row_data in table_data:
                            print(f"  {row_data}")
                else:
                    print("No tables found or error in table extraction.")

            # Clean up the dummy DOCX
            if os.path.exists(sample_docx_path):
                os.remove(sample_docx_path)
                print(f"\nCleaned up dummy DOCX: {sample_docx_path}")

        except Exception as e:
            print(f"Could not run DocxParser example: {e}")
            print("Ensure python-docx is installed and you have rights to create files here.")
    else:
        print("python-docx is not installed. Skipping DocxParser example.")
